import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

from docx import Document
from docx.shared import Inches as DocxInches, Pt as DocxPt, RGBColor as DocxRGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUTPUT_DIR = r"C:\Users\saisa\Documents\Work\Projects\Personal\AgenticBrowser\samples"
os.makedirs(OUTPUT_DIR, exist_ok=True)

# -------------------------------------------------------------
# 1. Financial Budget Spreadsheet (.xlsx)
# -------------------------------------------------------------
def create_budget_spreadsheet():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "FY2026 Budget Overview"
    ws.views.sheetView[0].showGridLines = True

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Calibri", size=16, bold=True, color="1F497D")
    section_font = Font(name="Calibri", size=12, bold=True, color="1F497D")
    total_font = Font(name="Calibri", size=11, bold=True)
    
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    subtotal_fill = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")
    net_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    double_bottom_border = Border(
        top=Side(style='thin', color='000000'),
        bottom=Side(style='double', color='000000')
    )

    # Title
    ws["B2"] = "Startup Operating Budget & Forecast (FY2026)"
    ws["B2"].font = title_font

    headers = ["Category / Line Item", "Q1 Actual", "Q2 Budget", "Q3 Budget", "Q4 Budget", "Full Year Total"]
    for col_idx, header in enumerate(headers, start=2):
        cell = ws.cell(row=4, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center" if col_idx > 2 else "left")

    # Revenue
    ws["B5"] = "REVENUE"
    ws["B5"].font = section_font
    rev_data = [
        ("SaaS Subscriptions", 45000, 60000, 85000, 120000),
        ("Enterprise Services", 15000, 20000, 25000, 30000),
        ("Partner Commissions", 5000, 7500, 10000, 15000)
    ]
    curr_row = 6
    for item, q1, q2, q3, q4 in rev_data:
        ws.cell(row=curr_row, column=2, value=item).border = thin_border
        for i, val in enumerate([q1, q2, q3, q4], start=3):
            c = ws.cell(row=curr_row, column=i, value=val)
            c.number_format = "$#,##0"
            c.border = thin_border
        total_cell = ws.cell(row=curr_row, column=7, value=f"=SUM(C{curr_row}:F{curr_row})")
        total_cell.number_format = "$#,##0"
        total_cell.border = thin_border
        curr_row += 1

    # Revenue Total
    rev_total_row = curr_row
    ws.cell(row=rev_total_row, column=2, value="Total Revenue").font = total_font
    ws.cell(row=rev_total_row, column=2).fill = subtotal_fill
    for col in range(3, 7):
        col_letter = get_column_letter(col)
        c = ws.cell(row=rev_total_row, column=col, value=f"=SUM({col_letter}6:{col_letter}{rev_total_row-1})")
        c.font = total_font
        c.fill = subtotal_fill
        c.number_format = "$#,##0"
    c_tot = ws.cell(row=rev_total_row, column=7, value=f"=SUM(G6:G{rev_total_row-1})")
    c_tot.font = total_font
    c_tot.fill = subtotal_fill
    c_tot.number_format = "$#,##0"

    # Expenses
    curr_row += 2
    ws.cell(row=curr_row, column=2, value="OPERATING EXPENSES").font = section_font
    curr_row += 1
    exp_start_row = curr_row
    exp_data = [
        ("Engineering & Product", 25000, 30000, 35000, 40000),
        ("Sales & Marketing", 18000, 22000, 28000, 32000),
        ("Cloud Infrastructure (AWS/GCP)", 6000, 8000, 10000, 12000),
        ("General & Administrative", 8000, 9000, 9500, 10000)
    ]
    for item, q1, q2, q3, q4 in exp_data:
        ws.cell(row=curr_row, column=2, value=item).border = thin_border
        for i, val in enumerate([q1, q2, q3, q4], start=3):
            c = ws.cell(row=curr_row, column=i, value=val)
            c.number_format = "$#,##0"
            c.border = thin_border
        total_cell = ws.cell(row=curr_row, column=7, value=f"=SUM(C{curr_row}:F{curr_row})")
        total_cell.number_format = "$#,##0"
        total_cell.border = thin_border
        curr_row += 1

    # Expense Total
    exp_total_row = curr_row
    ws.cell(row=exp_total_row, column=2, value="Total Expenses").font = total_font
    ws.cell(row=exp_total_row, column=2).fill = subtotal_fill
    for col in range(3, 7):
        col_letter = get_column_letter(col)
        c = ws.cell(row=exp_total_row, column=col, value=f"=SUM({col_letter}{exp_start_row}:{col_letter}{exp_total_row-1})")
        c.font = total_font
        c.fill = subtotal_fill
        c.number_format = "$#,##0"
    c_tot_exp = ws.cell(row=exp_total_row, column=7, value=f"=SUM(G{exp_start_row}:G{exp_total_row-1})")
    c_tot_exp.font = total_font
    c_tot_exp.fill = subtotal_fill
    c_tot_exp.number_format = "$#,##0"

    # Net Income / Profit
    curr_row += 2
    net_row = curr_row
    ws.cell(row=net_row, column=2, value="NET OPERATING INCOME").font = total_font
    ws.cell(row=net_row, column=2).fill = net_fill
    ws.cell(row=net_row, column=2).border = double_bottom_border
    for col in range(3, 8):
        col_letter = get_column_letter(col)
        c = ws.cell(row=net_row, column=col, value=f"={col_letter}{rev_total_row}-{col_letter}{exp_total_row}")
        c.font = total_font
        c.fill = net_fill
        c.number_format = "$#,##0"
        c.border = double_bottom_border

    # Adjust column widths
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 32
    for col in ["C", "D", "E", "F", "G"]:
        ws.column_dimensions[col].width = 16

    path = os.path.join(OUTPUT_DIR, "Financial_Budget_FY2026.xlsx")
    wb.save(path)
    print(f"Created: {path}")

# -------------------------------------------------------------
# 2. Pitch Deck Presentation (.pptx)
# -------------------------------------------------------------
def create_pitch_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    DARK_BLUE = RGBColor(26, 38, 57)
    ACCENT_BLUE = RGBColor(41, 128, 185)
    GRAY_TEXT = RGBColor(100, 110, 120)
    DARK_TEXT = RGBColor(30, 41, 59)
    WHITE = RGBColor(255, 255, 255)

    def add_header(slide, title_text, category="PITCH DECK"):
        tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11.5), Inches(1.2))
        tf = tb.text_frame
        tf.word_wrap = True
        p_sub = tf.paragraphs[0]
        p_sub.text = category.upper()
        p_sub.font.size = Pt(10)
        p_sub.font.bold = True
        p_sub.font.color.rgb = ACCENT_BLUE

        p_main = tf.add_paragraph()
        p_main.text = title_text
        p_main.font.size = Pt(26)
        p_main.font.bold = True
        p_main.font.color.rgb = DARK_TEXT

    # Slide 1: Title Slide (Dark Theme)
    s1 = prs.slides.add_slide(blank_layout)
    bg = s1.shapes.add_shape(1, 0, 0, Inches(13.333), Inches(7.5)) # Rectangle
    bg.fill.solid()
    bg.fill.fore_color.rgb = DARK_BLUE
    bg.line.color.rgb = DARK_BLUE

    tb1 = s1.shapes.add_textbox(Inches(1.5), Inches(2.2), Inches(10.3), Inches(3.0))
    tf1 = tb1.text_frame
    p1 = tf1.paragraphs[0]
    p1.text = "NovaFlow AI"
    p1.font.size = Pt(48)
    p1.font.bold = True
    p1.font.color.rgb = WHITE

    p2 = tf1.add_paragraph()
    p2.text = "Next-Generation Autonomous Workflow Orchestration"
    p2.font.size = Pt(22)
    p2.font.color.rgb = RGBColor(180, 200, 220)

    p3 = tf1.add_paragraph()
    p3.text = "\nSeed Round Pitch Deck  •  Confidential  •  2026"
    p3.font.size = Pt(14)
    p3.font.color.rgb = RGBColor(130, 150, 170)

    # Slide 2: Problem & Solution
    s2 = prs.slides.add_slide(blank_layout)
    add_header(s2, "The Market Opportunity: Fragmented Enterprise Workflows", "The Problem & Solution")

    # Box 1: Problem
    b1 = s2.shapes.add_shape(1, Inches(0.8), Inches(2.0), Inches(5.6), Inches(4.5))
    b1.fill.solid()
    b1.fill.fore_color.rgb = RGBColor(248, 250, 252)
    b1.line.color.rgb = RGBColor(226, 232, 240)
    tf_b1 = b1.text_frame
    tf_b1.margin_left = Inches(0.4)
    tf_b1.margin_top = Inches(0.4)
    p = tf_b1.paragraphs[0]
    p.text = "🚨 The Problem"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = RGBColor(180, 40, 40)
    
    problems = [
        "Knowledge workers waste 4.2 hrs/day switching between disjointed SaaS tools.",
        "Traditional RPA tools are fragile and break with any minor UI update.",
        "Custom integration pipelines cost enterprises $250k+ annually in maintenance."
    ]
    for prob in problems:
        p = tf_b1.add_paragraph()
        p.text = f"• {prob}"
        p.font.size = Pt(13)
        p.font.color.rgb = DARK_TEXT

    # Box 2: Solution
    b2 = s2.shapes.add_shape(1, Inches(6.8), Inches(2.0), Inches(5.6), Inches(4.5))
    b2.fill.solid()
    b2.fill.fore_color.rgb = RGBColor(240, 249, 255)
    b2.line.color.rgb = RGBColor(186, 230, 253)
    tf_b2 = b2.text_frame
    tf_b2.margin_left = Inches(0.4)
    tf_b2.margin_top = Inches(0.4)
    p = tf_b2.paragraphs[0]
    p.text = "💡 The NovaFlow Solution"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = RGBColor(12, 74, 110)

    solutions = [
        "Self-Healing Browser Agents: Vision-guided navigation that adapts dynamically.",
        "Unified Control Plane: Natural language prompts orchestrated into multi-agent tasks.",
        "Zero Code Integration: Connects existing internal tools without complex API maintenance."
    ]
    for sol in solutions:
        p = tf_b2.add_paragraph()
        p.text = f"• {sol}"
        p.font.size = Pt(13)
        p.font.color.rgb = DARK_TEXT

    # Slide 3: Traction & Ask
    s3 = prs.slides.add_slide(blank_layout)
    add_header(s3, "Traction & Seed Financing Round", "Growth & The Ask")

    metrics = [
        ("350%", "YoY ARR Growth"),
        ("$1.2M", "Current ARR"),
        ("125%", "Net Dollar Retention"),
        ("42", "Enterprise Pilots")
    ]
    for i, (val, label) in enumerate(metrics):
        box = s3.shapes.add_shape(1, Inches(0.8 + i * 2.95), Inches(2.2), Inches(2.8), Inches(2.0))
        box.fill.solid()
        box.fill.fore_color.rgb = RGBColor(241, 245, 249)
        box.line.color.rgb = RGBColor(203, 213, 225)
        tf = box.text_frame
        tf.margin_top = Inches(0.4)
        p = tf.paragraphs[0]
        p.text = val
        p.font.size = Pt(32)
        p.font.bold = True
        p.font.color.rgb = ACCENT_BLUE
        p.alignment = PP_ALIGN.CENTER
        
        p2 = tf.add_paragraph()
        p2.text = label
        p2.font.size = Pt(13)
        p2.font.color.rgb = GRAY_TEXT
        p2.alignment = PP_ALIGN.CENTER

    # The Ask Section
    ask_box = s3.shapes.add_shape(1, Inches(0.8), Inches(4.7), Inches(11.65), Inches(1.8))
    ask_box.fill.solid()
    ask_box.fill.fore_color.rgb = DARK_BLUE
    ask_box.line.color.rgb = DARK_BLUE
    tf_ask = ask_box.text_frame
    tf_ask.margin_left = Inches(0.5)
    tf_ask.margin_top = Inches(0.3)
    p = tf_ask.paragraphs[0]
    p.text = "🎯 Raising $3.5M Seed Funding"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = WHITE
    
    p = tf_ask.add_paragraph()
    p.text = "Use of Funds: 60% Core AI Engineering & Safety, 25% Go-to-Market Expansion, 15% Strategic Partnerships."
    p.font.size = Pt(14)
    p.font.color.rgb = RGBColor(200, 215, 230)

    path = os.path.join(OUTPUT_DIR, "NovaFlow_Pitch_Deck.pptx")
    prs.save(path)
    print(f"Created: {path}")

# -------------------------------------------------------------
# 3. Project Proposal Document (.docx)
# -------------------------------------------------------------
def create_project_proposal():
    doc = Document()
    
    # Title
    title = doc.add_paragraph()
    r = title.add_run("Project Proposal: Enterprise Intelligent Automation Hub")
    r.font.size = DocxPt(22)
    r.font.bold = True
    r.font.color.rgb = DocxRGBColor(31, 73, 125)

    sub = doc.add_paragraph()
    r_sub = sub.add_run("Prepared for: Executive Leadership Team  |  Date: October 2026  |  Version: 1.0")
    r_sub.font.size = DocxPt(10)
    r_sub.font.italic = True
    r_sub.font.color.rgb = DocxRGBColor(120, 120, 120)

    doc.add_heading("1. Executive Summary", level=1)
    doc.add_paragraph(
        "This proposal outlines the initiative to deploy an autonomous agentic orchestration layer across core business "
        "workflows. By reducing repetitive manual data entry between legacy CRM, ERP, and communication channels, this project "
        "projects a 40% reduction in operational cycle time and an estimated $420,000 annual cost savings within 12 months."
    )

    doc.add_heading("2. Project Objectives & Scope", level=1)
    obj = [
        "Unify browser-based SaaS workflows into automated agent task queues.",
        "Implement human-in-the-loop review guardrails for compliance and data governance.",
        "Integrate automated document parsing (invoices, client intake, contract verification)."
    ]
    for o in obj:
        doc.add_paragraph(o, style='List Bullet')

    doc.add_heading("3. Implementation Roadmap", level=1)
    table = doc.add_table(rows=1, cols=4)
    table.style = 'Light Shading Accent 1'
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Phase"
    hdr_cells[1].text = "Key Deliverable"
    hdr_cells[2].text = "Duration"
    hdr_cells[3].text = "Success Metric"

    phases = [
        ("Phase 1: Discovery & Architecture", "Infrastructure audit & agent sandbox", "Weeks 1-3", "Architecture approved"),
        ("Phase 2: Core Pilot Deployment", "Automate finance & intake workflows", "Weeks 4-8", ">95% Task accuracy"),
        ("Phase 3: Org-wide Rollout", "Full team onboarding & monitoring dashboard", "Weeks 9-12", "100+ active automated runs/wk")
    ]
    for phase, deliv, dur, metric in phases:
        row_cells = table.add_row().cells
        row_cells[0].text = phase
        row_cells[1].text = deliv
        row_cells[2].text = dur
        row_cells[3].text = metric

    doc.add_heading("4. Budget & Resource Allocation", level=1)
    doc.add_paragraph(
        "The estimated total project investment is $85,000, encompassing software licensing, cloud compute infrastructure, "
        "and specialized implementation consulting. Projected payback period is 4.5 months post-deployment."
    )

    path = os.path.join(OUTPUT_DIR, "Project_Proposal_Enterprise_Automation.docx")
    doc.save(path)
    print(f"Created: {path}")

if __name__ == "__main__":
    create_budget_spreadsheet()
    create_pitch_deck()
    create_project_proposal()
    print("ALL_GENERATED_SUCCESSFULLY")
